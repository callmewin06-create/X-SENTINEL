"""Insert a BODMAS appendix in a new copy; preserve every other package part."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'output/ieee_2026_10_16/X_SENTINEL_IEEE_1col.docx'
OUT = ROOT / 'output/bodmas_2026_10_23/X_SENTINEL_IEEE_1col_BODMAS_2026_10_23.docx'
QA = ROOT / 'outputs/bodmas_2026_10_23'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def text(node):
    return ''.join(node.xpath('.//w:t/text()', namespaces=NS))


def replace_span(node, old, new):
    original = text(node)
    assert original.count(old) == 1
    start, end = original.index(old), original.index(old) + len(old)
    offset, started = 0, False
    for t in node.xpath('.//w:t', namespaces=NS):
        value = t.text or ''
        stop = offset + len(value)
        if stop > start and offset < end:
            a, b = max(0, start - offset), min(len(value), end - offset)
            t.text = value[:a] + (new if not started else '') + value[b:]
            t.set(qn('xml:space'), 'preserve')
            started = True
        offset = stop
    assert text(node) == original.replace(old, new)


def add_table(doc, block):
    rows = block['rows']
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = table._tbl.tblPr.find(qn('w:tblW'))
    width.set(qn('w:type'), 'dxa')
    width.set(qn('w:w'), str(round(sum(block['widths']) * 1440)))
    for col, size in zip(table.columns, block['widths']):
        col.width = Inches(size)
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + name)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            edge.set(qn('w:' + key), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for name, value in [('top', '65'), ('bottom', '65'), ('left', '80'), ('right', '80')]:
        edge = OxmlElement('w:' + name)
        edge.set(qn('w:w'), value)
        edge.set(qn('w:type'), 'dxa')
        margins.append(edge)
    table._tbl.tblPr.append(margins)
    for i, values in enumerate(rows):
        row = table.rows[0] if i == 0 else table.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
        for cell, value, size in zip(row.cells, values, block['widths']):
            cell.width = Inches(size)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'F1F5F9')
                cell._tc.get_or_add_tcPr().append(shade)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9)
            run.font.bold = i == 0
            run.font.color.rgb = RGBColor(0, 0, 0)


def build_appendix(blocks):
    doc = Document()
    for block in blocks:
        kind = block['kind']
        if kind == 'table':
            add_table(doc, block)
            continue
        style = {'heading1': 'Heading 1', 'heading2': 'Heading 2', 'caption': 'Caption'}.get(kind, 'Normal')
        p = doc.add_paragraph(style=style)
        pf = p.paragraph_format
        pf.line_spacing = 1
        pf.space_after = Pt(4)
        pf.first_line_indent = Inches(.14 if kind == 'p' else 0)
        pf.keep_together = kind != 'p'
        pf.keep_with_next = kind in ('heading1', 'heading2', 'caption')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if kind == 'p' else WD_ALIGN_PARAGRAPH.CENTER if kind in ('heading1', 'caption') else WD_ALIGN_PARAGRAPH.LEFT
        if kind == 'heading1':
            pf.page_break_before = True
            pf.space_before = Pt(0)
            pf.space_after = Pt(8)
        elif kind == 'heading2':
            pf.space_before = Pt(6)
            if block['text'] == 'D Timing boundaries':
                # Keep the short timing table intact on the second appendix page.
                pf.page_break_before = True
        run = p.add_run(block['text'])
        run.font.name = 'Times New Roman'
        run.font.size = Pt(8.5 if kind == 'caption' else 10)
        run.font.bold = False
        run.font.italic = kind == 'heading2'
        run.font.color.rgb = RGBColor(0, 0, 0)
    return [deepcopy(el) for el in doc._element.body if el.tag != qn('w:sectPr')]


def main():
    blocks = json.loads((QA / 'report_blocks.json').read_text(encoding='utf8'))
    before_sha = sha(SOURCE.read_bytes())
    with ZipFile(SOURCE) as src:
        root = etree.fromstring(src.read('word/document.xml'))
        body = root.find('w:body', NS)
        before = [deepcopy(el) for el in body]
        intro = next(el for el in body if 'these other datasets are background references' in text(el))
        replace_span(intro,
            'This study trains and evaluates only LightGBM models; these other datasets are background references and are not additional experiments',
            'This study trains and evaluates only LightGBM models; SOREL-20M is background, while Appendix C adds a static BODMAS check with frozen EMBER-trained models')
        conclusion = next(el for el in body if text(el).startswith('The software includes a Vietnamese Streamlit'))
        replace_span(conclusion,
            'Appendix B records this deployment evidence and its timing boundaries.',
            'Appendix B records the historical deployment evidence and its timing boundaries. Appendix C adds member-supplied static BODMAS results and their verification limits.')
        changed = [i for i, (old, new) in enumerate(zip(before, body))
                   if etree.tostring(old, method='c14n', exclusive=True) != etree.tostring(new, method='c14n', exclusive=True)]
        assert len(changed) == 2, changed
        section = body.find('w:sectPr', NS)
        for element in build_appendix(blocks):
            section.addprevious(element)
        updated = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        OUT.parent.mkdir(parents=True, exist_ok=True)
        with ZipFile(OUT, 'w') as dest:
            for item in src.infolist():
                dest.writestr(item, updated if item.filename == 'word/document.xml' else src.read(item.filename))
    with ZipFile(SOURCE) as src, ZipFile(OUT) as dest:
        assert src.namelist() == dest.namelist()
        changed_parts = [name for name in src.namelist() if src.read(name) != dest.read(name)]
        assert changed_parts == ['word/document.xml']
    assert sha(SOURCE.read_bytes()) == before_sha
    proof = {'source': str(SOURCE), 'source_sha256': before_sha, 'output': str(OUT),
             'output_sha256': sha(OUT.read_bytes()), 'source_preserved': True,
             'changed_package_parts': changed_parts, 'changed_original_body_indices': changed,
             'existing_images_styles_references_tables_preserved': True,
             'new_appendix': 'C', 'new_tables': 2, 'new_inference': False}
    (QA / 'document_preservation.json').write_text(json.dumps(proof, indent=2) + '\n', encoding='utf8')
    print(json.dumps(proof, indent=2))


if __name__ == '__main__':
    main()
