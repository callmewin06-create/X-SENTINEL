"""Author the revised research manuscript in Word from verified evidence only."""
from pathlib import Path
import json,csv,re,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('outputs/revision_2026_10_13');OUT=Path('output/revision_2026_10_13');OUT.mkdir(parents=True,exist_ok=True)
S=json.loads((ROOT/'summary.json').read_text(encoding='utf-8'))
A=json.loads(Path('output/teacher_questions_2026_10_12/current_audit.json').read_text(encoding='utf-8'))
B=json.loads((ROOT/'binary_validation/validation.json').read_text(encoding='utf-8'))
REFS=json.loads((OUT/'references/catalog.json').read_text(encoding='utf-8'))
REFNUM={x['key']:i+1 for i,x in enumerate(REFS)}
doc=Document();section=doc.sections[0]
doc.settings.odd_and_even_pages_header_footer=False
section.different_first_page_header_footer=False
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=section.bottom_margin=Cm(1.9);section.left_margin=section.right_margin=Cm(2.0)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
    style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Normal'].font.bold=False
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.bold=False
for node in list(doc.styles['Title'].element.xpath('./w:pPr/w:pBdr')):
    node.getparent().remove(node)
for node in doc.styles.element.xpath('.//w:rFonts'):
    for attr in list(node.attrib):
        if 'theme' in attr.lower():del node.attrib[attr]
    if node.get(qn('w:ascii')) is None:node.set(qn('w:ascii'),'Times New Roman')
    if node.get(qn('w:hAnsi')) is None:node.set(qn('w:hAnsi'),'Times New Roman')
for name,size in [('Heading 1',13),('Heading 2',11.5)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(11);doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(10);doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.keep_with_next=True
doc.core_properties.title='X SENTINEL Cross View Signals for Backdoor Suspicion in PE Malware Classifiers'
doc.core_properties.author='Đặng Lý Quân; Trần Đặng Hữu Thắng; Lại Vĩnh Phú; Phạm Nguyễn Hoàng Phúc'
doc.core_properties.subject='Revised research manuscript aligned with frozen LightGBM artifacts and teacher feedback'
markdown=[];table_n=figure_n=0;bookmarks={};bookmark_id=0;citation_used=set()

def field(p,code,result):
    run=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
    text=OxmlElement('w:instrText');text.set(qn('xml:space'),'preserve');text.text=' '+code+' ';run._r.append(text)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');run._r.append(sep)
    p.add_run(str(result));end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');p.add_run()._r.append(end)

def text_cites(p,text):
    for i,part in enumerate(re.split(r'(\[ref\d{2}(?:,ref\d{2})*\]|\{\{\w+\}\})',text)):
        if part.startswith('[ref'):
            keys=part[1:-1].split(',');citation_used.update(keys);p.add_run('['+', '.join(str(REFNUM[k]) for k in keys)+']')
        elif part.startswith('{{'):
            key=part[2:-2];kind,num=bookmarks[key];p.add_run(kind+' ');field(p,'REF '+key+' \\h',num)
        else:p.add_run(part)

def p(text):
    para=doc.add_paragraph();text_cites(para,text);markdown.append(text)
    para.paragraph_format.widow_control=True;return para

def h(text,level=1):doc.add_heading(text,level);markdown.append('#'*(level+1)+' '+text)

def caption(key,label,text,num):
    global bookmark_id
    bookmarks[key]=(label,num);para=doc.add_paragraph(style='Caption');para.add_run(label+' ')
    bookmark_id+=1;start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_id));start.set(qn('w:name'),key);para._p.append(start)
    field(para,'SEQ '+label+' \\* ARABIC',num)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id));para._p.append(end)
    para.add_run('  '+text)

def table(key,title,headers,rows,widths=None):
    global table_n
    table_n+=1;caption(key,'Table',title,table_n)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [6.65/len(headers)]*len(headers)
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for cell,width in zip(t.rows[0].cells,widths):cell.width=Inches(width)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,width,text in zip(cells,widths,row):cell.width=Inches(width);cell.text=str(text)
    for i,row in enumerate(t.rows):
        props=row._tr.get_or_add_trPr();no_split=OxmlElement('w:cantSplit');props.append(no_split)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
            tcPr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EAF0' if i==0 else ('F6F7F9' if i%2==0 else 'FFFFFF'));tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                n=OxmlElement('w:'+edge);n.set(qn('w:w'),'75');n.set(qn('w:type'),'dxa');margins.append(n)
            tcPr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(9);run.bold=(i==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    markdown.append('**'+title+'**\n\n|'+'|'.join(headers)+'|\n|'+'|'.join(['---']*len(headers))+'|\n'+'\n'.join('|'+'|'.join(map(str,r))+'|' for r in rows))

def fig(key,title,path,width=6.35):
    global figure_n
    figure_n+=1;caption(key,'Figure',title,figure_n)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.add_run().add_picture(str(path),width=Inches(width))
    markdown.append('!['+title+']('+str(path)+')')

def eq(text):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    om=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=text;mr.append(mt);om.append(mr);para._p.append(om)
    markdown.append('`'+text+'`')

def pct(x):return '—' if x is None else f'{100*x:.2f}'
def val(x):return '—' if x is None else f'{x:.3f}'
def cohort(source,dataset,name,method):return next(x for x in S[source] if x['dataset']==dataset and x['cohort']==name and x['method']==method)

title=doc.add_paragraph('X SENTINEL Cross View Signals for Backdoor Suspicion in PE Malware Classifiers',style='Title')
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
for line in ('Đặng Lý Quân  •  Trần Đặng Hữu Thắng  •  Lại Vĩnh Phú  •  Phạm Nguyễn Hoàng Phúc',
    'FPT University, Ho Chi Minh City, Vietnam','danglyquan@gmail.com  •  callmewin06@gmail.com',
    'vinhphu2020@gmail.com  •  tqtphamnguyenhoangphuc@gmail.com'):
    para=doc.add_paragraph(line);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(10)
markdown.append('# '+doc.core_properties.title)
h('Abstract')
e18=cohort('original_cohorts','EMBER2018','restricted_confirmed','X-F');e24=cohort('original_cohorts','EMBER2024','restricted_confirmed','X-F')
f18=cohort('fresh_cohorts','EMBER2018','restricted_confirmed','X_primary_full');f24=cohort('fresh_cohorts','EMBER2024','restricted_confirmed','X_primary_full')
p(f'We investigate per-input backdoor suspicion in static PE malware classifiers using feature-view attribution and prediction signals. X-SENTINEL combines the maximum view share of absolute SHAP contributions with either a gated Behavioral SHAP signal or the positive probability gap between a clean Behavioral classifier and a suspicious main classifier. We evaluate frozen LightGBM models on EMBER2018 and the Win32, Win64 and .NET subset of EMBER2024. Each dataset has 27 poisoning configurations, and all configurations remain reported. Structural restricted edits and edits spanning three views are analyzed separately because neither has been verified as a realizable malware binary transformation. In the independently confirmed restricted cohort, full X-SENTINEL detects {pct(e18["recall"])}% of successful eligible attacks across 12 EMBER2018 runs and {pct(e24["recall"])}% across three EMBER2024 runs. On additional previously unused local source IDs, corresponding recall is {pct(f18["recall"])}% and {pct(f24["recall"])}%. An additional clean full-model shadow baseline detects substantially more attacks under richer access. Filtering the small number of failed PE parses changes some thresholds but does not establish dataset contamination as the cause of detection limitations. These results support a calibrated heuristic with explicit misses and resource assumptions; they do not establish universal backdoor detection or a future false-positive guarantee.')
p('Keywords: PE malware classification; backdoor poisoning; LightGBM; SHAP; attribution views; inference-time suspicion; calibration.')

h('1 Introduction')
p('Static malware classifiers summarize executable files as feature vectors and predict whether an input is malicious. EMBER provides a public setting for studying Windows PE classification, while EMBER2024 expands feature extraction and evaluation coverage. BODMAS and SOREL-20M provide additional settings that emphasize scale and temporal evaluation. This study trains and evaluates only LightGBM models; these other datasets are background references and are not additional experiments. [ref01,ref02,ref03,ref04,ref06]')
p('A poisoning attack changes training examples so that a model learns an unintended association. Backdoor attacks seek a target prediction when a trigger is present while retaining acceptable ordinary accuracy. Prior malware work demonstrates explanation-guided poisoning and selective triggering. Such results motivate inspection beyond the final malware label, but they do not guarantee that an arbitrary vector edit can be applied to an executable without changing its operation. [ref12,ref13,ref23]')
p('We ask whether feature-view concentration and disagreement can help identify individual suspicious predictions from a possibly compromised classifier. The output is an alert alongside the malware probability. A low alert score is not evidence that a file is safe, and one alert does not prove that the deployed model contains a backdoor. The study concerns controlled poisoning artifacts rather than incident attribution on an unknown production model.')
p('The contribution is an explicit implementation and paired evaluation of these signals under two resource settings, with locked benign calibration, complete attack records, component analysis and source-level uncertainty checks. We also report failures, compare an additional clean shadow baseline and examine the difference between vector edits and edits to a benign executable. The results constrain the claimed benefit rather than assuming that the proposed combination must win.')
h('1 1 Research questions',2)
p('RQ1: At separately calibrated thresholds, which successful attacks are detected by M3, M4 and their combination, and which are missed? RQ2: How do reduced and full resource settings compare on the same inputs and poisoned models, after separating restricted edits from edits spanning three views? RQ3: How stable are these observations under parse-flag sensitivity, previously unused local source IDs and a bounded adaptive optimizer? The former question about M5 is explicitly unanswered and moved to future work. Renumbering in this research manuscript does not assert that the course rubric has been changed.')

h('2 Related work')
h('2 1 Representations and evaluation of malware classifiers',2)
p('Feature-based PE classification is one approach among several. EMBER and EMBER2024 publish extractor-based representations, while MalConv studies raw executable bytes and Saxe and Berlin study a learned representation from binary program features. DREBIN concerns Android and is cited only as an example of a different platform and feature regime. Results from these settings cannot be transferred directly to the PE view map used here. [ref01,ref02,ref05,ref29,ref30]')
p('Gradient boosting and random forests provide historical context for tree classifiers. The current implementation uses LightGBM for the clean main, poisoned main and view models; it does not contain an SVM, neural malware classifier or random-forest experiment. SHAP and tree-specific attribution work support feature contribution analysis. LIME provides a related explanation approach but is not used in this detector. Explanations describe model behavior and do not establish executable semantics or causal compromise. [ref06,ref07,ref08,ref09,ref10,ref11,ref61]')
p('Security evaluation can be distorted by data reuse, temporal mismatch and an inappropriate claim boundary. TESSERACT and work revisiting TRANSCEND motivate reporting time and partition roles, while guidance on machine learning in computer security motivates explicit evaluation assumptions. Our fixed-source comparisons do not isolate a causal effect of dataset age, extractor version or software family. [ref14,ref15,ref17]')
h('2 2 Poisoning and evasion are different experiments',2)
p('BadNets, neural trojan attacks, hidden triggers, label-consistent poisoning, Poison Frogs and Sleeper Agent cover different training-time threat models. Backdoor surveys organize this literature, but assumptions such as access to training data, clean reference data or network activations vary across methods. Image-classifier evaluation offers useful warnings about scenario dependence without constituting a PE comparison. [ref31,ref32,ref33,ref34,ref35,ref36,ref37,ref58]')
p('Adversarial input modification is distinct from poisoning a model. General robustness studies, black-box attacks and adaptive evaluation explain why an attacker should be allowed to react to a defense and why failure of one search is not proof of robustness. The relation between learned features and adversarial examples does not justify a semantic invariant for imported functions in a PE file. [ref18,ref19,ref20,ref21,ref22,ref59]')
p('Malware evasion work investigates API-sequence attacks, reinforcement learning, byte-level perturbations and functionality-preserving optimization. These methods reinforce the need to distinguish feature space from executable modification. We do not claim that replacing hash buckets corresponds to adding a valid import or that the present poisoning trigger preserves a malicious payload. [ref23,ref24,ref25,ref26,ref27,ref28]')
h('2 3 Defenses and calibration',2)
p('STRIP studies perturbation-based prediction behavior. Spectral signatures, activation clustering and SPECTRE inspect representations; Neural Cleanse, Fine-Pruning and ABS analyze or modify models under their own assumptions. Anti-Backdoor Learning, attention distillation, decoupled training and Beatrix offer further defense strategies. These are literature comparisons, not baselines reimplemented for LightGBM here. Our STRIP adaptation perturbs numerical feature vectors and therefore differs from a physically valid PE transformation. [ref39,ref40,ref41,ref42,ref43,ref44,ref45,ref46,ref47,ref48,ref49]')
p('Finite calibration and distribution-free risk literature motivate distinguishing empirical calibration behavior from a deployment guarantee. The threshold here is a specific empirical order statistic; the manuscript does not claim to implement conformal risk control. Bootstrap and correlated ROC analysis also motivate preserving shared input dependence. Wilson intervals describe individual proportions, while AUROC and precision–recall literature motivate naming the evaluated positive class precisely. [ref50,ref51,ref52,ref53,ref54,ref55,ref56,ref57]')

h('3 Method and resource assumptions')
h('3 1 Attacker and defender access',2)
p('The controlled attacker modifies selected benign fit vectors while keeping their label at zero. Poison counts are 1,000, 2,000 and 4,000, corresponding to 0.5%, 1% and 2% of the complete 200,000-vector fit pool, or 1%, 2% and 4% of its benign half. The main model architecture and training settings are fixed. The target is malware-to-benign prediction after adding the trigger. Success and stealth are measured screening conditions, not invariants promised for every input.')
table('access','Resources available during scoring',['Setting','Main model','Benign reference and calibration','Additional clean models'],[
    ['Reduced','Suspicious LightGBM','500 reference and 2,000 calibration','None'],
    ['Full','Same suspicious model','Same source IDs','Three clean view models; M4 uses Behavioral'],
    ['Shadow comparison','Same suspicious model','Same source IDs','Additional clean 500-round full model']], [1.0,1.3,2.0,2.35])
p('The three view models were prepared from 200,000 clean labeled fit records, using 150 boosting rounds. Full therefore has substantially more information than a defender holding only a suspicious model and a few hundred benign samples. The shadow comparison is given a clean 500-round full model trained on that same fit pool. Reduced does not receive this model. Scoring primary detectors does not read final labels, poison IDs or the trigger manifest; the evaluator uses them to construct experiments.')
fig('scorepath','Implemented score path and separate benign threshold selection',ROOT/'analysis/figures/implemented_score_path.png')
p('{{scorepath}} shows the primary core. STRIP is an independently scored comparator, and the main probability remains available separately from the alert. Reduced combines M3 and M4R; full combines M3 and M4F.')
h('3 2 Versioned feature views',2)
table('views','Complete disjoint feature views',['Dataset','Dim.','Structural','Behavioral','Metadata','Behavioral index interval'],[
    ['EMBER2018',2381,357,1408,616,'[943, 2351)'],['EMBER2024',2568,467,1411,690,'[994, 2405)']],[1.25,.65,.85,.85,.8,2.25])
p('Behavioral denotes static import and export features. It is not a runtime trace, and it is not assumed immutable or monotone under all executable modifications. Metadata includes byte and string information. The versioned schema and model dimension are checked together; a V2 model cannot be evaluated as a V3 model. {{views}} replaces the incorrect V3 mapping in the source manuscript.')
h('3 3 Components and fusion',2)
p('Let p(x) be the main probability of malware and φj(x) its native LightGBM contribution to the raw margin, excluding the bias term. Let A be the total absolute contribution and Nv the signed contribution sum in view v. Contributions are interpreted as model evidence rather than verified program behavior. The epsilon is 10⁻⁹; the attribution ratios return zero when A is at most epsilon.')
eq('A = Σj |φj|     Nv = Σj∈v φj')
eq('TADR = maxj |φj| / A')
eq('M3 = maxv (Σj∈v |φj| / A)')
eq('M4R = (1 − p(x)) max(NBehavioral, 0) / A')
eq('M4F = max(0, pBehavioral(xBehavioral) − p(x))')
p('M3 measures concentration of absolute SHAP mass in a view. It is not the signed difference NBehavioral − NStructural. M4R is large when the main classifier favors benign while its Behavioral contribution is positive. M4F is a positive probability discrepancy. These are different signals, and changing to a polarity-contrast formula would require separate results. Full versus reduced changes both the information supplied and the M4 definition; the present comparison does not isolate these effects causally.')
eq('Rref(s) = #{r in benign reference : score(r) ≤ s} / 500')
eq('XR = (Rref(M3) + Rref(M4R)) / 2')
eq('XF = (Rref(M3) + Rref(M4F)) / 2')
p('Reference ranks use the right empirical CDF and ties are retained. Rank values and fusion values are not probabilities that an input contains a backdoor. Each component and combination has its own threshold for each poisoned model. TADR is the local maximum-feature absolute-attribution comparator; no universal 0.65 threshold is used.')
h('3 4 STRIP and thresholds',2)
p('For each input, STRIP draws 50 benign reference vectors with replacement and forms numerical blends with alpha 0.5. Binary entropy uses the prediction of each blend, and the score is one minus the mean entropy. The random sequence depends on the input content and seed so that batch order does not change the score. This is mean entropy, not entropy of the average probability. [ref39]')
eq('H(p) = −p log₂ p − (1 − p) log₂(1 − p)')
eq('STRIP(x) = 1 − (1/50) Σk H(p(0.5x + 0.5rk))')
p('Sort the 2,000 benign calibration scores ascending and use the 1,980th score as tau for the 1% target. An input is flagged only if score > tau. This ensures at most 20 strict exceedances in that calibration sample and can yield fewer because of ties. It does not guarantee an unseen benign FPR of at most 1%. Thresholds and reference ranks are never selected from the final attack outcomes.')
h('3 5 Decision policy and unanswered mechanism',2)
p('The detector exposes malware probability and suspicion separately. For evaluation, flagged successful attacks are assumed to be stopped by a simulated triage policy. Post-defense ASR is therefore a policy statistic, not a measured recovery of the classifier or a real-world incident outcome. M5 is disabled in all primary bundles and has no primary ablation. Hashed feature buckets cannot establish the presence of a specific API; causal reconstruction remains future work.')

h('4 Experimental protocol')
h('4 1 Source and partition provenance',2)
p('EMBER2024 vectors come from the joyce8/EMBER2024 dataset on Hugging Face, pinned to revision 3d23efef7c0f0b702c5024400cfff4c3744a3832. The six selected archives are Dot_Net_train.zip, Dot_Net_test.zip, Win32_train.zip, Win32_test.zip, Win64_train.zip and Win64_test.zip. Extraction code is from FutureComputing4AI/EMBER2024, pinned to commit 0ef753e81d98bf209f71b03cd331dfc190b5b54d. Dataset source and extractor source are separate provenance records. No APK, ELF or PDF archive is selected in this matrix. [ref02]')
role_rows=[]
for name in ('fit','selection','development','confirmation','reference','calibration','final'):
    r=A['roles'][name];role_rows.append([name,f"{r['n']:,}",f"{r['types']['Win32']:,}",f"{r['types']['Win64']:,}",f"{r['types']['Dot_Net']:,}",r['is_pe_zero'],f"{r['is_pe_zero_benign']}/{r['is_pe_zero_malware']}"])
table('roles','EMBER2024 source composition and parse flags by role',['Role','n','Win32','Win64','.NET','is_pe 0','Benign malware'],role_rows,[1.25,.85,.95,.9,.8,.8,1.1])
p('Both datasets use fit 200,000; selection, development and confirmation 10,000 each; reference 500 benign; calibration 2,000 benign; and final 20,000 balanced records. Balanced roles have equal benign and malware counts. Source IDs are disjoint by role, and all final models within a dataset use the same source order. Source-format labels and parser success flags are different observations. EMBER2024 contains 90 is_pe=0 records across these roles, including two calibration records and three final benign records. Without their original bytes, a cause of parse failure is not established.')
h('4 2 Models and attack screening',2)
table('settings','Locked training and attack settings',['Setting','Value'],[
    ['Main and clean full model','500 rounds; 64 leaves; learning rate 0.05'],['Clean view models','150 rounds; same clean fit source IDs'],
    ['Sampling','feature_fraction 0.9; bagging_fraction 0.9'],['Model seeds','17, 29, 43'],['Poison rates of all fit','0.5%, 1%, 2%'],
    ['Inference','4 CPU threads; malware cutoff 0.5'],['Concentrated restricted','2 Structural columns'],['Spread restricted','10 Structural columns'],
    ['Cross view stress','8 Structural + 8 Behavioral + 8 Metadata columns'],['Structural permitted pool','13 columns per schema; vector feasibility only']],[2.3,4.35])
p('Two selector candidates were evaluated during development, and legacy_rare was locked for the final matrix. Strong-and-stealth screening requires at least 1,000 eligible malware, eligible ASR at least 50%, paired gain at least 20 percentage points and strict clean-accuracy loss below 0.5 percentage points. Development and separate confirmation must both pass for a configuration to enter the confirmed cohort. These rules select viable attack conditions; they do not prove that the detector works. All 27 final configurations per dataset remain available regardless of confirmation status.')
p('The restricted cohort contains concentrated and spread Structural vector edits. Cross edits alter all three views and are placed in a separate stress analysis. Whitelisted numerical edits have not been verified as transformations of malicious executables. The term clean-label describes retained training labels within the vector experiment; it does not certify the validity of edited bytes.')
h('4 3 Evaluation groups and denominators',2)
table('denominators','Attack evaluation groups',['Group','Membership or denominator'],[
    ['All triggered malware','All 10,000 final malware after applying the fixed trigger'],
    ['Eligible malware','Clean full model predicts malware on the original input'],
    ['Successful eligible','Eligible and poisoned main predicts benign on the triggered input'],
    ['Paired new evasion','Successful eligible and clean full model still predicts malware after the same trigger'],
    ['Benign FPR','Flagged original final benign / 10,000 benign'],
    ['ASR before defense','Successful eligible / eligible'],['Recall','Flagged successful eligible / successful eligible'],
    ['Post defense ASR','Unflagged successful eligible / eligible']],[1.65,5.0])
p('AUROC is reported separately for all triggered malware, successful eligible attacks and paired new evasions against the same benign negative set. These are different positive classes. A successful eligible input can still reflect direct evasion if the trigger also fools the clean model. The stricter paired group helps examine this possibility. Means are calculated over runs, giving equal weight to each run. Mean post-defense ASR is not obtained by multiplying mean ASR by one minus mean recall.')
h('4 4 Supplementary validation and bounded adaptive evaluation',2)
p('After review, we selected a fresh local validation pool of 1,000 benign and 1,000 malware for each dataset. Selection uses source metadata and deterministic rules before predictions, excluding all known primary roles, local pilot IDs and attack-development partition IDs. EMBER2018 uses unused official test records; EMBER2024 uses unused labeled test records located through the source audit and the pinned raw archives. One new EMBER2024 record retains is_pe=0. External handover observation history is unavailable, so independence is stated relative to known local use. Frozen triggers, models, cohorts and thresholds are reused without outcome-based reselection.')
p('New shadow and inverse-STRIP variants were proposed after observing the old final pool, so their old-final results are exploratory. The shadow score is max(0, pclean − pmain), with a threshold from the original benign calibration. Inverse STRIP uses the original mean entropy H and its own benign-calibrated threshold; it is not the same detector with a changed name. We validate both on the unused local pool without retuning alpha, K, fusion weights or thresholds.')
p('The bounded adaptive study uses the first 100 fresh malware source IDs per originally confirmed restricted run, selected without conditioning on outcomes. The attacker sees probability and reduced fusion score and knows the formula. The goal is pmain < 0.5 and XR ≤ its old tau. The original trigger coordinates remain fixed; proposals change remaining permitted Structural coordinates using values observed in benign selection data. Each input has at most 32 sequential queries, including the initial query. This is a fixed coordinate search in vector space, not a physically verified PE attack or an exhaustive robustness assessment. Preliminary runs that accidentally admitted V2 Metadata whitelist entries were withdrawn and excluded from the reported Structural-only experiment.')

h('5 Results')
h('5 1 Which attack configurations support the principal comparison',2)
table('cohorts','Confirmation status by attack family',['Dataset','Concentrated confirmed','Spread confirmed','Cross confirmed','All confirmed'],[
    ['EMBER2018','3 / 9','9 / 9','0 / 9','12 / 27'],['EMBER2024','0 / 9','3 / 9','9 / 9','12 / 27']],[1.25,1.5,1.3,1.25,1.35])
p('The two confirmed groups have different composition. The restricted EMBER2024 group is only three spread runs, all at 2% poisoning. Comparing its mean with 12 restricted EMBER2018 runs cannot isolate a dataset effect. Cross stress accounts for nine of the 12 confirmed EMBER2024 runs and is therefore separated from the restricted headline. {{cohorts}} includes the unsuccessful confirmation conditions rather than treating the number 12 as universal coverage.')
rows=[]
for dataset in ('EMBER2018','EMBER2024'):
    for method in ('TADR','STRIP','M3','M4-R','M4-F','X-R','X-F'):
        r=cohort('original_cohorts',dataset,'restricted_confirmed',method);rows.append([dataset.replace('EMBER',''),method,r['runs'],pct(r['recall']),pct(r['fpr']),pct(r['post_asr']),val(r['auc_successful'])])
table('restricted','Restricted confirmed cohort at original thresholds',['Dataset','Method','Runs','Recall %','FPR %','Post ASR %','AUC success'],rows,[.85,.7,.5,1.0,1.0,1.1,1.5])
p('{{restricted}} reports operating-point performance and successful-eligible AUROC on the same positive class. On EMBER2018, combining M3 and reduced M4 improves mean recall over M3 alone. On EMBER2024, a combination can rank differently from a component at the calibrated operating point, and full resources do not imply reliable coverage. Different final FPR values also prevent a blanket conclusion of superiority at matched deployment cost.')
for dataset in ('EMBER2018','EMBER2024'):
    fig('recall'+dataset,'Component and fusion recall for '+dataset+' restricted confirmed runs',ROOT/'analysis/figures'/f'{dataset}_restricted_recall.png')
h('5 2 Cross view stress and complete configuration coverage',2)
cross=[]
for method in ('TADR','M3','M4-R','M4-F','X-R','X-F'):
    r=cohort('original_cohorts','EMBER2024','cross_confirmed',method);cross.append([method,r['runs'],pct(r['recall']),pct(r['fpr']),pct(r['asr_before']),pct(r['post_asr'])])
table('cross','EMBER2024 confirmed cross stress with its own scope',['Method','Runs','Recall %','FPR %','ASR before %','Post ASR %'],cross,[1.0,.65,1.2,1.1,1.35,1.35])
p('There is no confirmed EMBER2018 cross cohort. Results for weak cross configurations still appear in the full configuration appendix. Spreading numerical changes across views is a stress condition for view concentration, but these measurements do not prove that an attacker can realize the original 24-column trigger in a functioning PE binary. TADR can detect some dispersed-trigger cases; the previous exactly-zero general claim is removed.')
h('5 3 Paired differences and complementary detections',2)
cluster_rows=[]
for r in S['cluster']:
    if r['cohort']!='restricted_confirmed':continue
    lo,hi=r['recall']['ci95'];cluster_rows.append([r['dataset'].replace('EMBER',''),r['first'].replace('M3_view_mass','M3').replace('M4_prob_gap','M4-F').replace('X_primary_reduced','X-R'),r['second'].replace('X_primary_reduced','X-R').replace('X_primary_full','X-F'),r['runs'],pct(r['recall']['difference']),f'{pct(lo)} to {pct(hi)}'])
table('paired','Source ID clustered recall contrasts on fixed models',['Dataset','First','Second','Runs','Difference pp','95% interval pp'],cluster_rows,[.8,1.0,1.0,.5,1.3,2.05])
p('The 1,000 bootstrap draws resample each source ID once and retain that ID across runs, so repeated inputs are not treated as independent observations. {{paired}} is conditional on the observed models and chosen cohorts; it does not provide uncertainty over future seeds or attack selectors. The small EMBER2024 restricted cohort supports a descriptive conclusion for these three models, not a general factorial claim.')
catch=list(csv.DictReader((ROOT/'analysis/catch_miss.csv').open()))
catch_rows=[]
for dataset in ('EMBER2018','EMBER2024'):
    selected=[r for r in catch if r['dataset']==dataset and r['confirmed']=='True' and r['family']!='cross_3view_stress' and r['group']=='successful_eligible']
    catch_rows.append([dataset.replace('EMBER',''),len(selected)]+[sum(int(r[k]) for r in selected) for k in ('n','both','first_only','second_only','neither')])
table('catch','TADR and M4F flags on successful eligible restricted attacks',['Dataset','Runs','Occurrences','Both','TADR only','M4F only','Neither'],catch_rows,[.8,.45,1.2,.85,1.1,1.1,1.15])
p('Counts in {{catch}} are source/model occurrences because the same malware source can appear in multiple runs. They show which inputs each comparator misses or catches; they are not unique independent samples. Per-run catch/miss and rank correlations remain in the machine-readable artifact so that illustrative examples are not substituted for complete evidence.')
h('5 4 Parse flag sensitivity and calibration',2)
sensrows=[]
for r in S['sensitivity']:
    if r['dataset']=='EMBER2024':sensrows.append([r['method'].replace('M3_view_mass','M3').replace('M4_shap_behavioral_conflict','M4-R').replace('M4_prob_gap','M4-F').replace('X_primary_reduced','X-R').replace('X_primary_full','X-F'),r['runs'],val(r['max_abs_tau_delta']),pct(r['max_abs_recall_delta']),pct(r['max_abs_fpr_delta'])])
table('sensitivity','Largest observed changes across 27 EMBER2024 runs after filtering parse flags',['Method','Runs','Abs tau change','Recall change pp','FPR change pp'],sensrows,[1.1,.6,1.5,1.7,1.75])
p('Filtering the three final benign parse failures and filtering the two calibration parse failures are evaluated separately and jointly while holding the model, trigger and reference fixed. {{sensitivity}} reports the maximum absolute joint change across all runs rather than claiming that two records are necessarily harmless. The experiment does not retrain the fit pool and does not identify parse failures as the cause of limited recall. No non-PE archive contamination is observed in the selected source provenance.')
p('The two original calibration records have all-zero Behavioral input but nonzero Behavioral SHAP under an audited model: signed totals are approximately 0.3460 and 0.7536. This directly refutes the implication that zero input must yield zero attribution. All 90 original parse-failure records have zero feature tails from index 696 onward. Their missing information is a data-quality observation that requires contextual analysis, not a proof of a particular performance mechanism.')
p('Calibration scores were recomputed for every model and all old thresholds reproduced within 10⁻¹⁰. For a continuous IID order-statistic diagnostic, the exceedance probability has Beta(21,1980) uncertainty before considering the final benign count. Beta-binomial upper tails and calibration exceedance counts are provided per run. Discrete scores, ECDF ties, dependence across runs and possible train/test differences limit this diagnostic; a high observed FPR is not by itself causal proof of distribution shift.')
h('5 5 STRIP score direction and additional clean model comparison',2)
for dataset in ('EMBER2018','EMBER2024'):
    fig('entropy'+dataset,'Original benign and triggered mean entropy for '+dataset,ROOT/'analysis/figures'/f'{dataset}_strip_entropy.png')
p('The entropy plots distinguish the observed score direction from an explanation invented after the fact. In the old confirmed cohorts, mean triggered entropy is higher than benign entropy, so low AUROC under score 1−H is not explained by both groups having identically maximal entropy. Inverse STRIP changes the ranking direction, but AUROC improvement alone does not show high recall at its own benign-calibrated threshold.')
baseline_rows=[]
for dataset in ('EMBER2018','EMBER2024'):
    for method in ('shadow','strip_inverse'):
        old=cohort('supplement_old_final',dataset,'restricted_confirmed',method)
        fresh=cohort('fresh_cohorts',dataset,'restricted_confirmed','STRIP_inverse' if method=='strip_inverse' else method)
        baseline_rows.append([dataset.replace('EMBER',''),method.replace('_',' '),old['runs'],pct(old['recall']),pct(old['fpr']),pct(fresh['recall']),pct(fresh['fpr'])])
table('baselines','New baselines on restricted cohorts with thresholds from original benign calibration',['Dataset','Method','Runs','Old recall %','Old FPR %','Fresh recall %','Fresh FPR %'],baseline_rows,[.7,1.25,.45,1.05,1.05,1.1,1.05])
p('The clean shadow comparator is strong under explicitly richer access. {{baselines}} does not assign its capability to reduced. A clean full model at 500 rounds also has a different model budget from the 150-round Behavioral classifier; the contrast does not isolate an architectural cause. Where the shadow comparator wins, the contribution must be framed as a study of attribution signals under access restrictions, rather than a claim that adding view models produces the best defense.')
h('5 6 Previously unused source validation and bounded adaptation',2)
freshrows=[]
for dataset in ('EMBER2018','EMBER2024'):
    for method in ('TADR','M3_view_mass','M4_shap_behavioral_conflict','M4_prob_gap','X_primary_reduced','X_primary_full'):
        r=cohort('fresh_cohorts',dataset,'restricted_confirmed',method)
        label={'M3_view_mass':'M3','M4_shap_behavioral_conflict':'M4-R','M4_prob_gap':'M4-F','X_primary_reduced':'X-R','X_primary_full':'X-F'}.get(method,method)
        freshrows.append([dataset.replace('EMBER',''),label,r['runs'],pct(r['asr_before']),pct(r['recall']),pct(r['fpr']),pct(r['post_asr'])])
table('fresh','Fresh local source IDs under frozen original attack selection',['Dataset','Method','Runs','ASR before %','Recall %','FPR %','Post ASR %'],freshrows,[.8,.7,.5,1.2,1.1,1.1,1.25])
p('The fresh pool has 1,000 benign and 1,000 malware per dataset, so FPR resolution per run is 0.1 percentage points. Successful-eligible sample counts vary by model and are saved explicitly. The original cohort membership is retained instead of selecting new favorable attacks from this pool. {{fresh}} is supplementary validation on new local source IDs, not a new independent training replication.')
table('adaptive','Bounded Structural vector optimizer on fresh malware',['Dataset','Runs','Eligible occurrences','Initial joint','Final joint','New joint','Max queries'],[
    [r['dataset'].replace('EMBER',''),r['runs'],r['eligible_occurrences'],r['start_joint_occurrences'],r['end_joint_occurrences'],r['new_joint_occurrences'],r['max_queries']] for r in S['adaptive']],[.8,.45,1.45,1.0,1.0,1.0,.95])
p('Joint success requires both a benign main prediction and no reduced alert. {{adaptive}} separates already-unflagged attacks from newly achieved joint successes within the same eligible set. Occurrence counts repeat source IDs across models; they must not be read as independent attacks. The optimizer retains original trigger coordinates and stays within the 13 Structural columns. Limited new success would constrain this search only; it would not establish resilience to import padding, new poisoning selectors or stronger adaptive search.')
h('5 7 Cost of the classifier and detector',2)
table('latency','Mean of run level warm per input timing summaries in milliseconds',['Dataset','Scope','Runs','Mean ms','Median ms','P95 ms'],[
    [r['dataset'].replace('EMBER',''),r['scope'].replace('_',' '),r['runs'],val(r['mean_ms']),val(r['median_ms']),val(r['p95_ms'])] for r in S['timing']],[.8,1.95,.5,1.1,1.1,1.2])
p('Timing uses an AMD Ryzen 7 8745H processor with 8 physical cores, 16 logical processors and approximately 15.31 GiB of reported physical memory, under Windows and four inference threads. The same 100 inputs are used within each model, after warm-up, with alternating reversed evaluation order. The component shortcut is checked against the original detector scores. Classifier, reduced core, full core, STRIP and complete suites are timed separately. The averages of per-run medians and p95 values in {{latency}} are not pooled distribution percentiles. Extraction, model loading, API transport and database writes are excluded; ordinary system background load is not controlled. A 10 ms core target is assessed against its defined scope; no sub-millisecond end-to-end claim is made.')
h('5 8 Benign executable validation',2)
binary_rows=[]
for variant in B['variants']:
    for version,d in variant['feature_differences'].items():
        binary_rows.append([variant['variant'].replace('_',' '),'V2' if 'v2' in version else 'V3','4 / 4',d['changed_n'],d['changed_by_view']['structural'],d['changed_by_view']['behavioral'],d['changed_by_view']['metadata'],len(d['outside_restricted'])])
table('binary','Feature changes after byte edits to one self built benign executable',['Edit','Ver.','Cases','Changed','Structural','Behavioral','Metadata','Outside pool'],binary_rows,[1.5,.5,.6,.6,.8,.85,.8,.9])
p('We compiled one reviewed unsigned .NET program that prints n²+7, without file or network operations. Copies received two COFF timestamp substitutions and a benign overlay. Each copy preserved exit code, stdout and stderr on four inputs, including an invalid input. Only these self-built benign fixtures were executed. The V3 reader uses checksum-pinned upstream parsing and its binary/raw vector parity is checked. The V2 modern-LIEF reader remains experimental because identical-binary parity with legacy LIEF 0.9 has not been established.')
p('Even timestamp edits change byte and string-derived features outside the intended Structural coordinate, as shown in {{binary}}. Preserved outputs on four benign cases demonstrate a small prototype, not preservation of every behavior. This experiment does not realize the original vector triggers, validate any malicious payload or demonstrate a successful binary backdoor. It gives a concrete reason to keep these evidence levels separate.')

h('6 Discussion and limitations')
p('The original arithmetic can be internally consistent while the narrative is wrong. The revised method now states the formulas that actually produced the stored results. Absolute view mass is a concentration heuristic, and probability discrepancy is an additional signal under clean-model access. Neither supplies a proof of semantic inconsistency or universal backdoor presence. The failure cases and the strong shadow comparison reduce the scope of an appropriate contribution.')
p('An apparent dataset difference is confounded by view schemas, source periods, attack viability and confirmed-family composition. EMBER2024 has only three confirmed restricted runs at one rate. Source and parse-flag audits remove a particular ambiguity in the report; they do not rule out every data or representation issue. Filtering calibration/test flags without retraining is a sensitivity experiment, not a repaired full training protocol.')
p('Ablations show that fusion can change operating-point coverage, but a few seeds and fixed thresholds do not prove stable performance over future models. Equal reference-rank fusion is retained because it is the locked method; no weights are optimized using final outcomes. Richer baselines must be reported with their information and model budgets. The bounded adaptive optimizer covers only one declared numerical attack family and is not a complete adversarial evaluation.')
p('References and bibliography identities have been checked against registered or primary metadata. That verification is not a claim that every cited paper has been independently replicated. Claims are limited to the particular concepts for which sources are cited. Raw malware functionality, operational alert handling, rare/packed benign cohorts, M5 effectiveness and deployment security remain outside the measured research evidence.')

h('7 Reproducibility and conclusion')
p('The research artifact retains all original models, triggers, thresholds and result files. Revision outputs have a separate dated namespace with code hashes, source IDs, per-input scores, calibration caches, diagnostics and new validation protocols. A reproduction package distinguishes recreating tables from saved scores from retraining models on raw data. No public repository release, DOI or external submission is implied by this local package.')
p('The software includes a Streamlit demonstration, an API adapter and stored analysis history. These demonstrate an interface to the trained detector rather than measured authentication, backup recovery or end-to-end scanning efficacy. Current automated checks pass 58 tests. The new V3 reader is a read-only research component validated on self-built benign files; it is not an asserted production upload route.')
p('X-SENTINEL provides measurable per-input suspicion signals with substantial misses and explicit access requirements. On the evaluated frozen models, combining attribution and view-related evidence helps in some settings but does not establish a generally strongest defense. Additional clean-model access can support a stronger simple comparator. The next research steps are independently specified binary transformations, broader adaptive tests, resource-controlled clean-model comparisons and prospective calibration checks. M5 remains an unanswered extension.')

doc.add_page_break();h('References')
for i,r in enumerate(REFS,1):
    text=f"[{i}] "+', '.join(r['authors'])+f', “{r["title"]},” '+r['venue']+f', {r["year"]}. '+r['url']
    para=doc.add_paragraph(text);para.paragraph_format.space_after=Pt(5);para.paragraph_format.keep_together=True
    para.paragraph_format.left_indent=Inches(.25);para.paragraph_format.first_line_indent=Inches(-.25)
    for run in para.runs:run.font.size=Pt(10)
    markdown.append(text)
if citation_used!=set(REFNUM):raise ValueError('Uncited bibliography entries: '+str(set(REFNUM)-citation_used))

doc.add_page_break();h('Appendix A Complete frozen configuration results')
p('Each row retains an original final configuration. C, S and X denote concentrated Structural, spread Structural and cross-view stress. Poison rate is a percentage of all fit records. Confirmed means both development and separate confirmation pass the locked screen. All values below are from the original final pool and thresholds. The artifact provides all seven methods, counts, three AUROC definitions and catch/miss records per run.')
raw=list(csv.DictReader((ROOT/'analysis/cell_metrics.csv').open()));appendix=[]
for dataset in ('EMBER2018','EMBER2024'):
    for r in [x for x in raw if x['dataset']==dataset and x['method']=='X-F']:
        rr=next(x for x in raw if all(x[k]==r[k] for k in ('dataset','seed','family','rate')) and x['method']=='X-R')
        family={'concentrated_structural':'C','spread_structural':'S','cross_3view_stress':'X'}[r['family']]
        appendix.append([dataset.replace('EMBER',''),r['seed'],family,f"{float(r['rate'])*100:g}",'Yes' if r['confirmed']=='True' else 'No',pct(float(r['asr_before'])),pct(float(rr['recall'])),pct(float(r['recall'])),pct(float(r['fpr']))])
table('allconfigs','All 54 configurations with no removal of weak attacks',['Dataset','Seed','Family','Rate %','Confirmed','ASR %','XR recall %','XF recall %','XF FPR %'],appendix,[.8,.5,.55,.55,.75,.75,.9,.9,.9])
p('Source audit, four parse-flag conditions, the original versus fresh comparisons, rank correlations and withdrawn preliminary adaptive implementation are recorded in the revision artifact. Corresponding-author designation, ORCID, target venue and course requirements for the earlier M5 question require author or teacher confirmation and are not invented in this manuscript.')

header=section.header.paragraphs[0];header.text='X SENTINEL  |  Revised research manuscript';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:run.font.size=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(footer,'PAGE',1)
settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
path=OUT/'X_SENTINEL_Report_Revised.docx';doc.save(path)
(OUT/'manuscript.md').write_text('\n\n'.join(markdown),encoding='utf-8')
(OUT/'report_build.json').write_text(json.dumps({'docx':str(path),'tables':table_n,'figures':figure_n,'references':len(REFS),
    'citations_used':len(citation_used),'summary_sha256':hashlib.sha256((ROOT/'summary.json').read_bytes()).hexdigest(),
    'source_report_preserved':True,'format':'Word single-column working manuscript; native equations and automatic captions/references'},indent=2),encoding='utf-8')
print('Word manuscript created',path,'tables',table_n,'figures',figure_n,'references',len(REFS))
