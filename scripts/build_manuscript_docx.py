import os
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import latex2mathml.converter
from lxml import etree
import win32com.client

# Initialize MathML to OMML transformer
xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt_doc = etree.parse(xslt_path)
transform = etree.XSLT(xslt_doc)

def latex_to_omml(latex_str, display='inline'):
    clean_tex = latex_str.strip()
    try:
        mml_str = latex2mathml.converter.convert(clean_tex, display=display)
        mml_tree = etree.fromstring(mml_str.encode('utf-8'))
        omml_tree = transform(mml_tree)
        omml_bytes = etree.tostring(omml_tree, encoding='utf-8')
        return parse_xml(omml_bytes)
    except Exception as e:
        return None

def create_document():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = False
        
        # Professional Header
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h_run = p_hdr.add_run("IEEE TRANSACTIONS ON INFORMATION FORENSICS AND SECURITY (REVISION MANUSCRIPT)")
        h_run.font.name = "Times New Roman"
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = RGBColor(140, 140, 140)
        
        # Professional Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = p_ft.add_run("X-SENTINEL: Feature-View Attribution Concentration | Page ")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(120, 120, 120)
        
        fldSimple = OxmlElement('w:fldSimple')
        fldSimple.set(qn('w:instr'), 'PAGE')
        r_page = OxmlElement('w:r')
        t_page = OxmlElement('w:t')
        t_page.text = "1"
        r_page.append(t_page)
        fldSimple.append(r_page)
        p_ft._p.append(fldSimple)
        
    return doc

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="10" w:space="0" w:color="0F172A"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="10" w:space="0" w:color="0F172A"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_table_margins(table, top=45, bottom=45, left=65, right=65):
    tblPr = table._tbl.tblPr
    tblCellMar = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )
    tblPr.append(tblCellMar)

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def add_formatted_runs(paragraph, text, base_font_size=10, base_color=None):
    # Step 1: Split by inline math ($...$) first so math is NEVER swallowed by bold/italic
    math_chunks = re.split(r'(\$.*?\$)', text)
    
    for chunk in math_chunks:
        if not chunk:
            continue
        if chunk.startswith('$') and chunk.endswith('$') and len(chunk) >= 2:
            tex_content = chunk[1:-1].strip()
            omml_elem = latex_to_omml(tex_content, display='inline')
            if omml_elem is not None:
                paragraph._p.append(omml_elem)
                continue
            # Fallback
            run = paragraph.add_run(tex_content)
            run.font.name = "Cambria Math"
            run.font.size = Pt(base_font_size)
            run.italic = True
            continue
            
        # Step 2: For regular text, split by markdown formatting: ***, **, *, `
        style_pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)')
        parts = style_pattern.split(chunk)
        
        for part in parts:
            if not part:
                continue
            bold = False
            italic = False
            clean_text = part
            font_name = "Times New Roman"
            font_size = base_font_size
            color = base_color
            
            if part.startswith('***') and part.endswith('***') and len(part) >= 6:
                bold = True
                italic = True
                clean_text = part[3:-3]
            elif part.startswith('**') and part.endswith('**') and len(part) >= 4:
                bold = True
                clean_text = part[2:-2]
            elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
                italic = True
                clean_text = part[1:-1]
            elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
                font_name = "Consolas"
                font_size = base_font_size - 0.5
                clean_text = part[1:-1]
                color = RGBColor(30, 41, 59)
                
            run = paragraph.add_run(clean_text)
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.bold = bold
            run.italic = italic
            if color:
                run.font.color.rgb = color

def render_table(doc, table_lines):
    if len(table_lines) < 2:
        return
        
    rows_data = []
    for l in table_lines:
        inner = l.strip()[1:-1]
        cols = [c.strip() for c in inner.split('|')]
        rows_data.append(cols)
        
    headers = rows_data[0]
    data_rows = rows_data[2:] if len(rows_data) > 1 and all(set(c).issubset({'-', ':', ' '}) for c in rows_data[1]) else rows_data[1:]
    
    num_cols = len(headers)
    num_rows = len(data_rows) + 1
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    set_table_borders(table)
    set_table_margins(table, top=45, bottom=45, left=65, right=65)
    
    # Header Row
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    # Header bottom border & text
    for c_idx, title in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        set_cell_background(cell, "F1F5F9")
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="475569"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(tcBorders)
        
        p = cell.paragraphs[0]
        # Align column 0 left if descriptive, else center
        if c_idx == 0 and any(w in title.lower() for w in ['dataset', 'cohort', 'setting', 'fixture', 'role', 'parameter']):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        add_formatted_runs(p, title, base_font_size=8.5, base_color=RGBColor(15, 23, 42))
        
    # Data Rows
    for r_idx, r_data in enumerate(data_rows):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        
        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell_text = r_data[c_idx] if c_idx < len(r_data) else ""
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
                
            p = cell.paragraphs[0]
            clean_val = cell_text.replace('%', '').replace('+', '').replace('-', '').strip()
            is_numeric = False
            try:
                float(clean_val)
                is_numeric = True
            except ValueError:
                is_numeric = False
                
            if is_numeric or cell_text.endswith('%') or '/' in cell_text or clean_val.lower() in ('yes', 'no', 'true', 'false'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_formatted_runs(p, cell_text, base_font_size=7.5 if num_cols >= 10 else 8.0)
            
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    doc = create_document()
    lines = content.splitlines()
    
    table_lines = []
    
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i].strip()
        
        # Check if table row
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            i += 1
            continue
        elif table_lines:
            render_table(doc, table_lines)
            table_lines = []
            
        if not line:
            i += 1
            continue
            
        # Title
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line[2:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue
            
        # Authors
        if line.startswith('**Đặng Lý Quân') or line.startswith('**Author') or 'Department of Information Assurance' in line:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs(p, line, base_font_size=10.5)
            i += 1
            continue
            
        if line.startswith('`{danglyquan') or line.startswith('{danglyquan'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(16)
            add_formatted_runs(p, line, base_font_size=9.5, base_color=RGBColor(80, 80, 80))
            i += 1
            continue

        if line == '---':
            i += 1
            continue
            
        # Section Heading (## ...)
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            
            title_text = line[3:].strip()
            if re.match(r'^(I|II|III|IV|V|VI|VII|VIII|IX|X)\.\s+', title_text) or title_text.startswith('APPENDIX') or title_text.startswith('REFERENCES'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(title_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(16, 44, 87)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(title_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue
            
        # Subsection Heading (### ...)
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            
            sub_text = line[4:].strip()
            if sub_text.lower() == 'abstract':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run("ABSTRACT")
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(16, 44, 87)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(sub_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.italic = True
                run.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue
            
        # Sub-subsection Heading (#### ...)
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            sub_sub_text = line[5:].strip()
            run = p.add_run(sub_sub_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.italic = True
            i += 1
            continue
            
        # Display Math Equation ($$...$$)
        if line.startswith('$$') and line.endswith('$$'):
            eq_text = line[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            
            omml_elem = latex_to_omml(eq_text, display='block')
            if omml_elem is not None:
                p._p.append(omml_elem)
            else:
                run = p.add_run(eq_text)
                run.font.name = "Cambria Math"
                run.font.size = Pt(10.5)
                run.italic = True
            i += 1
            continue
            
        # Table Caption / Title
        if line.startswith('**TABLE ') or line.startswith('TABLE '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_formatted_runs(p, line, base_font_size=9.5, base_color=RGBColor(16, 44, 87))
            i += 1
            continue

        # Bullet List
        if line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run_b = p.add_run("•  ")
            run_b.font.name = "Times New Roman"
            run_b.font.size = Pt(10)
            run_b.bold = True
            add_formatted_runs(p, line[2:].strip(), base_font_size=10)
            i += 1
            continue
            
        # Numbered List
        num_match = re.match(r'^(\d+)\.\s+(.*)', line)
        if num_match:
            num = num_match.group(1)
            text_body = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            is_ref = ('arXiv' in text_body or 'Proc.' in text_body or 'IEEE' in text_body or 'Anderson' in text_body or 'NeurIPS' in text_body or 'USENIX' in text_body)
            run_num = p.add_run(f"[{num}]  " if is_ref else f"{num}.  ")
            run_num.font.name = "Times New Roman"
            run_num.font.size = Pt(9.5 if is_ref else 10)
            run_num.bold = True
            add_formatted_runs(p, text_body, base_font_size=9.5 if is_ref else 10)
            i += 1
            continue
            
        # Body Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        
        if 'We investigate per-input backdoor suspicion' in line or 'All backdoor poisoning' in line:
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            add_formatted_runs(p, line, base_font_size=9.5)
        elif line.startswith('**Keywords**:'):
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(12)
            add_formatted_runs(p, line, base_font_size=9.5)
        else:
            add_formatted_runs(p, line, base_font_size=10)
            
        i += 1
        
    if table_lines:
        render_table(doc, table_lines)
        
    doc.save(docx_path)
    print(f"Generated docx: {docx_path}")

if __name__ == '__main__':
    md_file = r'G:\IAM301t\X-SENTINEL\docs\REVISED_MANUSCRIPT_IEEE_X_SENTINEL.md'
    docx_file = r'G:\IAM301t\X-SENTINEL\docs\REVISED_MANUSCRIPT_IEEE_X_SENTINEL_PERFECT.docx'
    parse_markdown_to_docx(md_file, docx_file)
    
    # Resave via Word native COM to guarantee 100% canonical formatting
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        wdoc = word.Documents.Open(docx_file, ConfirmConversions=False, ReadOnly=False)
        print("Opened in Word, saving canonical version...")
        final_docx = r'G:\IAM301t\X-SENTINEL\docs\REVISED_MANUSCRIPT_IEEE_X_SENTINEL.docx'
        wdoc.SaveAs2(final_docx, 16)
        
        final_pdf = r'G:\IAM301t\X-SENTINEL\docs\REVISED_MANUSCRIPT_IEEE_X_SENTINEL.pdf'
        wdoc.SaveAs2(final_pdf, 17)
        
        dl_docx = r'C:\Users\LENOVO\Downloads\REVISED_MANUSCRIPT_IEEE_X_SENTINEL.docx'
        wdoc.SaveAs2(dl_docx, 16)
        
        dl_pdf = r'C:\Users\LENOVO\Downloads\REVISED_MANUSCRIPT_IEEE_X_SENTINEL.pdf'
        wdoc.SaveAs2(dl_pdf, 17)
        
        print("Pages:", wdoc.ComputeStatistics(2))
        print("OMaths:", wdoc.OMaths.Count)
        print("Tables:", wdoc.Tables.Count)
        wdoc.Close(False)
        print("ALL EXPORTS COMPLETED PERFECTLY!")
    finally:
        word.Quit()
        
    shutil.copy2(final_docx, r'C:\Users\LENOVO\.gemini\antigravity\brain\1088f6f4-15f9-4d4f-8a3b-d4de37cb0eff\REVISED_MANUSCRIPT_IEEE_X_SENTINEL.docx')
    print("Artifact updated!")
