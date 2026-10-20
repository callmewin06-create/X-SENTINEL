"""Build the two BODMAS handoff documents from their reviewed Markdown sources."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output/bodmas_demo_2026_10_19"
SOURCES = {
    "plan": ROOT / "docs/plans/KE_HOACH_TRIEN_KHAI_VA_DEMO_2026_10_19.md",
    "guide": ROOT / "docs/HUONG_DAN_BODMAS_KALI_2026_10_19.md",
}
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^\s)]+\))")


def set_font(run, name="Arial", size=None, bold=None):
    run.font.name = name
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._r.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for key in list(fonts.attrib):
        if "theme" in key.lower():
            del fonts.attrib[key]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def hyperlink(paragraph, label, target):
    link = OxmlElement("w:hyperlink")
    rid = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rp = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "145A86")
    rp.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rp.append(underline)
    r.append(rp)
    t = OxmlElement("w:t")
    t.text = label
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def inline(paragraph, text, font_size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", 9.0 if font_size is None else min(font_size, 9.0))
        elif part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=font_size, bold=True)
        elif part.startswith("["):
            m = re.fullmatch(r"\[([^\]]+)\]\(([^\s)]+)\)", part)
            hyperlink(paragraph, m.group(1), m.group(2))
        else:
            for item in re.split(r"(https?://[^\s]+)", part):
                if item.startswith(("http://", "https://")):
                    url = item.rstrip(".,;)")
                    hyperlink(paragraph, url, url)
                    if len(url) < len(item):
                        set_font(paragraph.add_run(item[len(url):]), size=font_size)
                elif item:
                    set_font(paragraph.add_run(item), size=font_size)


def clean_heading(text):
    text = text.replace("`", "").replace("**", "")
    # Headings use words and numbers; technical punctuation remains in body text.
    text = re.sub(r"[^\w\s]", " ", text, flags=re.UNICODE)
    return re.sub(r"\s+", " ", text).strip()


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    section.footer_distance = Inches(0.3)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        for border in list(style.element.xpath("./w:pPr/w:pBdr")):
            border.getparent().remove(border)
        style.paragraph_format.widow_control = True
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.07
    normal.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(19)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    doc.styles["Title"].paragraph_format.keep_with_next = True
    for name, size in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    code = doc.styles.add_style("Shell command", 1)
    code.base_style = doc.styles["Normal"]
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.04)
    code.paragraph_format.line_spacing = 1.04
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.widow_control = True
    doc.styles["Footer"].font.name = "Arial"
    doc.styles["Footer"].font.size = Pt(9)
    doc.styles["Footer"].font.color.rgb = RGBColor(0, 0, 0)
    # Remove theme font and color residue from the styles used by Word.
    for node in doc.styles.element.xpath(".//w:rFonts"):
        for key in list(node.attrib):
            if "theme" in key.lower():
                del node.attrib[key]
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        for color in doc.styles[name].element.xpath(".//w:color"):
            color.attrib.clear()
            color.set(qn("w:val"), "000000")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Trang ")
    set_font(r, size=9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = "Nhóm X SENTINEL"
    doc.core_properties.subject = "Kế hoạch và hướng dẫn kiểm tra tĩnh BODMAS trên Kali Linux"


def cells(line):
    return [s.strip().replace(r"\|", "|") for s in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def add_table(doc, lines):
    parsed = [cells(s) for s in lines]
    if len(parsed) > 1 and all(re.fullmatch(r":?-+:?", x.replace(" ", "")) for x in parsed[1]):
        parsed.pop(1)
    n = len(parsed[0])
    if any(len(row) != n for row in parsed):
        raise ValueError("Malformed Markdown table")
    # Allocate more width to columns carrying explanations rather than short values.
    weights = []
    for col in range(n):
        sizes = sorted(len(row[col]) for row in parsed)
        typical = sizes[min(len(sizes) - 1, int(len(sizes) * 0.7))]
        weights.append(max(12, min(62, typical)) ** 0.7)
    widths = [7.1 * x / sum(weights) for x in weights]
    t = doc.add_table(rows=0, cols=n)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    t._tbl.tblPr.append(borders)
    for i, data in enumerate(parsed):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, (cell, text, width) in enumerate(zip(row.cells, data, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + edge)
                node.set(qn("w:w"), "90")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            tcpr.append(margins)
            if i == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DCE6EF")
                tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.03
            p.paragraph_format.keep_with_next = i == 0
            inline(p, text.replace("<br>", "\n"), font_size=9.7)
            if i == 0:
                for r in p.runs:
                    r.bold = True
            elif len(text) < 16 and not any(c.isalpha() for c in text):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = Pt(3)


def build(source):
    text = source.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    doc = Document()
    configure(doc)
    title = None
    code_blocks = 0
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            code_blocks += 1
            for k, value in enumerate(block):
                p = doc.add_paragraph(style="Shell command")
                p.paragraph_format.keep_with_next = len(block) <= 16 and k < len(block) - 1
                p.paragraph_format.keep_together = True
                set_font(p.add_run(value), "Consolas", 9)
            if block:
                p.paragraph_format.space_after = Pt(7)
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.+)$", line)
        if m:
            heading = clean_heading(m.group(2))
            if title is None and len(m.group(1)) == 1:
                title = heading
                p = doc.add_paragraph(heading, style="Title")
            else:
                p = doc.add_heading(heading, level=min(3, max(1, len(m.group(1)) - 1)))
            for r in p.runs:
                set_font(r)
            i += 1
            continue
        if line.lstrip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if re.fullmatch(r"\s*[-*_]{3,}\s*", line):
            i += 1
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        numbered = re.match(r"^(\s*)(\d+)[.)]\s+(.+)$", line)
        if bullet or numbered:
            match = bullet or numbered
            p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            p.paragraph_format.space_after = Pt(4)
            if len(match.group(1)):
                p.paragraph_format.left_indent = Inches(0.25 + 0.12 * min(3, len(match.group(1)) // 2))
            inline(p, match.group(2) if bullet else match.group(3))
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            inline(p, line[2:])
            i += 1
            continue
        parts = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|```|\s*[-*]\s|\s*\d+[.)]\s|\s*\|)", lines[i]):
            parts.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        inline(p, " ".join(parts))
    if title is None:
        raise ValueError("Source requires a Markdown document title")
    doc.core_properties.title = title
    dest = OUTPUT / (source.stem + ".docx")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return {
        "source": str(source),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "output": str(dest),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "code_blocks": code_blocks,
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--only", choices=("all", "plan", "guide"), default="all")
    a = p.parse_args()
    selected = list(SOURCES.values()) if a.only == "all" else [SOURCES[a.only]]
    for path in selected:
        if not path.is_file():
            raise FileNotFoundError(path)
    records = [build(path) for path in selected]
    qa = OUTPUT / "documents"
    qa.mkdir(parents=True, exist_ok=True)
    (qa / "authoring.json").write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(records, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
